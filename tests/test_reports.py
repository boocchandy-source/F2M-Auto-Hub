"""Report rendering and Excel export tests."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from f2m.models import (
    Appointment,
    Customer,
    FeedbackConcern,
    Notification,
    ServiceHistory,
    ServiceRequest,
    Visit,
    db,
)
from f2m.timeutil import now_str, today_iso


@pytest.fixture()
def report_data(app):
    today = today_iso()
    with app.app_context():
        customer = Customer(
            name="Report Customer",
            plate="RPT 100",
            model="Toyota Hilux",
            address="Quezon, Bukidnon",
            technician="Jun Reyes",
            created_at=now_str(),
        )
        db.session.add(customer)
        db.session.flush()
        db.session.add_all(
            [
                Visit(
                    visit_date=today,
                    record_no="1001",
                    plate="RPT 100",
                    model="Toyota Hilux",
                    customer_name="Report Customer",
                    address="Quezon, Bukidnon",
                    technician="Jun Reyes",
                ),
                Appointment(
                    customer_id=customer.id,
                    service="Wheel Alignment",
                    appointment_date=today,
                    appointment_time="10:30",
                    status="Scheduled",
                    created_at=now_str(),
                ),
                ServiceRequest(
                    customer_id=customer.id,
                    service="Tire Replacement",
                    request_date=today,
                    status="Completed",
                    technician="Jun Reyes",
                    notes="Two front tires.",
                    created_at=now_str(),
                ),
                ServiceHistory(
                    customer_id=customer.id,
                    service="Tire Replacement",
                    service_date=today,
                    technician="Jun Reyes",
                    status="Completed",
                    notes="Two front tires.",
                ),
                FeedbackConcern(
                    customer_id=customer.id,
                    concern="Customer requested a follow-up call.",
                    status="Pending",
                    created_at=now_str(),
                ),
                Notification(
                    customer_id=customer.id,
                    message="Your vehicle is ready for pickup.",
                    notification_date=now_str(),
                    status="Pending",
                ),
            ]
        )
        db.session.commit()
        return customer.id


def test_daily_report_renders(app, authed_client, report_data):
    response = authed_client.get(f"/reports?type=Daily&date={today_iso()}")
    assert response.status_code == 200
    assert b"DAILY REPORT" in response.data
    assert b"Report Customer" in response.data
    assert b"Download .xlsx" in response.data


@pytest.mark.parametrize("report_type", ["Daily", "Weekly", "Monthly"])
def test_excel_export_is_a_valid_workbook(app, authed_client, report_data, report_type):
    response = authed_client.get(
        f"/reports/export?type={report_type}&date={today_iso()}"
    )
    assert response.status_code == 200
    assert response.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["Content-Disposition"]
    assert response.data[:2] == b"PK"

    workbook = load_workbook(BytesIO(response.data))
    assert "SUMMARY" in workbook.sheetnames
    assert len(workbook.sheetnames) > 1

    expected_day = datetime.strptime(today_iso(), "%Y-%m-%d").strftime("%B %d, %Y")
    assert expected_day in workbook.sheetnames

    daily_sheet = workbook[expected_day]
    values = [
        cell
        for row in daily_sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    ]
    assert "Report Customer" in values

    requests_sheet = workbook["Service Requests"]
    request_values = [
        cell
        for row in requests_sheet.iter_rows(values_only=True)
        for cell in row
        if cell
    ]
    assert "Tire Replacement" in request_values

    summary = workbook["SUMMARY"]
    summary_values = [
        cell for row in summary.iter_rows(values_only=True) for cell in row if cell
    ]
    assert f"{report_type} Report" in summary_values or report_type in summary_values


def test_report_with_invalid_parameters_falls_back(app, authed_client):
    response = authed_client.get("/reports?type=Unknown&date=not-a-date")
    assert response.status_code == 200
    assert b"DAILY REPORT" in response.data


def test_export_rows_exist_for_records(app, authed_client, report_data):
    response = authed_client.get(f"/reports/export?type=Daily&date={today_iso()}")
    workbook = load_workbook(BytesIO(response.data))
    for sheet in (
        "Customer Records",
        "Appointments",
        "Service Requests",
        "Service History",
        "Feedback",
        "Notifications",
    ):
        assert sheet in workbook.sheetnames

    with app.app_context():
        assert db.session.scalar(select(Customer.id)) is not None
