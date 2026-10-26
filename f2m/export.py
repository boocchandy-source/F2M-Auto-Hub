"""Excel report export built with openpyxl."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .timeutil import now_local

BRAND = "D62828"
DARK = "202327"
MUTED = "747A82"
SOFT = "F4F5F7"
LINE = "D9DCE1"
WHITE = "FFFFFF"

THIN = Side(style="thin", color=LINE)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor=DARK)
BRAND_FILL = PatternFill("solid", fgColor=BRAND)
SOFT_FILL = PatternFill("solid", fgColor=SOFT)

TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=WHITE)
BRAND_FONT = Font(name="Calibri", size=15, bold=True, color=WHITE)
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=10, color="17191C")
BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="17191C")
MUTED_FONT = Font(name="Calibri", size=10, color=MUTED)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="top", wrap_text=True)
LEFT_CENTER = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _safe(value) -> str:
    if value is None:
        return ""
    text = ILLEGAL_CHARACTERS_RE.sub("", str(value))
    # Prevent Excel from interpreting imported text as a formula.
    if text.startswith("="):
        text = "'" + text
    return text


def _set_value(cell, value) -> None:
    cell.value = value if isinstance(value, (int, float)) else _safe(value)
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = LEFT_CENTER


def _merge_title(ws, row: int, text: str, last_column: int, fill, font, height: int = 28) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_column)
    cell = ws.cell(row=row, column=1, value=_safe(text))
    cell.fill = fill
    cell.font = font
    cell.alignment = CENTER
    ws.row_dimensions[row].height = height
    for column in range(1, last_column + 1):
        ws.cell(row=row, column=column).border = BORDER


def _write_headers(ws, row: int, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=index, value=_safe(header))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = CENTER
    ws.row_dimensions[row].height = 23


def _write_rows(ws, start_row: int, rows: list[list], widths: list[int] | None = None) -> int:
    for offset, values in enumerate(rows):
        for index, value in enumerate(values, start=1):
            _set_value(ws.cell(row=start_row + offset, column=index), value)
    if widths:
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
    return start_row + len(rows)


def _auto_widths(rows: list[list], minimum: int = 10, maximum: int = 55) -> list[int]:
    if not rows:
        return []
    widths = []
    for index in range(len(rows[0])):
        longest = max((len(_safe(row[index])) for row in rows), default=0)
        widths.append(max(minimum, min(maximum, longest + 3)))
    return widths


def _safe_sheet_title(name: str, used: set[str]) -> str:
    for character in "[]:*?/\\":
        name = name.replace(character, "_")
    name = (name.strip() or "Sheet")[:31]
    base = name
    counter = 2
    while name in used:
        suffix = f" ({counter})"
        name = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name)
    return name


def _date_label(iso_date: str) -> str:
    try:
        return datetime.strptime(iso_date, "%Y-%m-%d").strftime("%B %d, %Y")
    except ValueError:
        return iso_date


def _daily_sheet(wb: Workbook, visit_date: str, rows: list, used: set[str]):
    sheet_name = _safe_sheet_title(_date_label(visit_date), used)
    ws = wb.create_sheet(sheet_name)
    ws.sheet_view.showGridLines = False

    widths = [8, 18, 25, 36, 38, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    _merge_title(ws, 1, "F2M AUTO-HUB TIRES & MAGS TRADING", 6, BRAND_FILL, TITLE_FONT, 30)
    _merge_title(ws, 2, "Blgy 1A, Poblacion, Quezon, Bukidnon", 6, SOFT_FILL, MUTED_FONT, 22)
    ws.merge_cells("A5:F5")
    title = ws["A5"]
    title.value = "DAILY CUSTOMER RECORD"
    title.font = Font(name="Calibri", size=14, bold=True, color="17191C")
    title.alignment = CENTER
    ws.row_dimensions[5].height = 28
    _merge_title(ws, 6, _date_label(visit_date).upper(), 6, SOFT_FILL, BOLD_FONT, 22)

    headers = ["NO.", "PLATE #", "MODEL", "CUSTOMER NAME", "ADDRESS", "TECHNICIAN"]
    _write_headers(ws, 8, headers)

    if rows:
        for offset, visit in enumerate(rows, start=9):
            values = [
                visit.record_no or (offset - 8),
                visit.plate or "-",
                visit.model or "-",
                visit.customer_name or "NO NAME",
                visit.address or "-",
                visit.technician or "-",
            ]
            for index, value in enumerate(values, start=1):
                _set_value(ws.cell(row=offset, column=index), value)
            ws.row_dimensions[offset].height = 20
    else:
        ws.merge_cells(start_row=9, start_column=1, end_row=9, end_column=6)
        cell = ws.cell(row=9, column=1, value="No customer records for this date.")
        cell.font = MUTED_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    ws.freeze_panes = "A9"
    ws.print_title_rows = "8:8"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.35, right=0.35, top=0.45, bottom=0.45, header=0.2, footer=0.2)
    return ws


def _table_sheet(
    wb: Workbook,
    title: str,
    headers: list[str],
    rows: list[list],
    used: set[str],
    widths: list[int] | None = None,
):
    ws = wb.create_sheet(_safe_sheet_title(title, used))
    ws.sheet_view.showGridLines = False
    last_column = len(headers)
    _merge_title(ws, 1, title.upper(), last_column, BRAND_FILL, BRAND_FONT, 30)
    _write_headers(ws, 3, headers)
    end_row = _write_rows(ws, 4, rows, widths or _auto_widths(rows, maximum=50))
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(last_column)}{end_row - 1}"
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def _summary_sheet(
    wb: Workbook,
    report_type: str,
    start: date,
    end: date,
    counts: dict[str, int],
):
    ws = wb.active
    ws.title = "SUMMARY"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 42
    for column in "CDEF":
        ws.column_dimensions[column].width = 18

    _merge_title(ws, 1, "F2M AUTO-HUB TIRES & MAGS TRADING", 6, BRAND_FILL, TITLE_FONT, 32)
    _merge_title(ws, 2, "Blgy 1A, Poblacion, Quezon, Bukidnon", 6, SOFT_FILL, MUTED_FONT, 22)
    _merge_title(ws, 4, "REPORT SUMMARY", 6, HEADER_FILL, BRAND_FONT, 28)

    info = [
        ("Report Type", report_type.title()),
        ("Period Start", start.strftime("%B %d, %Y")),
        ("Period End", end.strftime("%B %d, %Y")),
        ("Generated", now_local().strftime("%B %d, %Y %I:%M %p")),
    ]
    row = 6
    for label, value in info:
        label_cell = ws.cell(row=row, column=1, value=label)
        label_cell.font = BOLD_FONT
        label_cell.fill = SOFT_FILL
        label_cell.border = BORDER
        value_cell = ws.cell(row=row, column=2, value=value)
        value_cell.font = BODY_FONT
        value_cell.border = BORDER
        value_cell.alignment = LEFT_CENTER
        row += 1

    _write_headers(ws, 11, ["CATEGORY", "COUNT", "INCLUDED IN EXPORT"])
    rows = [
        ["Daily Visits", counts["visits"], "Yes"],
        ["Customer Records", counts["customers"], "Yes"],
        ["Appointments", counts["appointments"], "Yes"],
        ["Service Requests", counts["requests"], "Yes"],
        ["Service History", counts["history"], "Yes"],
        ["Feedback / Concerns", counts["concerns"], "Yes"],
        ["Notifications", counts["notifications"], "Yes"],
    ]
    _write_rows(ws, 12, rows, widths=None)
    for offset in range(7):
        ws.row_dimensions[12 + offset].height = 20
    ws.freeze_panes = "A11"
    return ws


def build_report_workbook(
    report_type: str,
    start: date,
    end: date,
    visits: list,
    customers: list,
    appointments: list,
    requests: list,
    history: list,
    concerns: list,
    notifications: list,
) -> BytesIO:
    workbook = Workbook()
    workbook.properties.creator = "F2M Auto-Hub"
    workbook.properties.title = f"F2M Auto-Hub {report_type} Report"
    workbook.properties.description = (
        "Customer service and automotive service management report generated by F2M Auto-Hub."
    )

    counts = {
        "visits": len(visits),
        "customers": len(customers),
        "appointments": len(appointments),
        "requests": len(requests),
        "history": len(history),
        "concerns": len(concerns),
        "notifications": len(notifications),
    }
    _summary_sheet(workbook, report_type, start, end, counts)

    grouped: dict[str, list] = {}
    for visit in visits:
        grouped.setdefault(visit.visit_date, []).append(visit)
    if not grouped:
        grouped[start.isoformat()] = []
    used: set[str] = {"SUMMARY"}
    for visit_date in sorted(grouped):
        _daily_sheet(workbook, visit_date, grouped[visit_date], used)

    customer_rows = [
        [
            customer.name,
            customer.plate or "-",
            customer.model or "-",
            customer.address or "-",
            customer.technician or "-",
            (customer.created_at or "")[:10],
        ]
        for customer in customers
    ]
    _table_sheet(
        workbook,
        "Customer Records",
        ["NAME", "PLATE NUMBER", "VEHICLE MODEL", "ADDRESS", "TECHNICIAN", "REGISTERED DATE"],
        customer_rows,
        used,
        [26, 16, 24, 32, 22, 16],
    )

    appointment_rows = [
        [
            appointment.id,
            appointment.customer.name if appointment.customer else "Unknown",
            appointment.service,
            appointment.appointment_date,
            appointment.appointment_time,
            appointment.status,
        ]
        for appointment in appointments
    ]
    _table_sheet(
        workbook,
        "Appointments",
        ["ID", "CUSTOMER", "SERVICE", "DATE", "TIME", "STATUS"],
        appointment_rows,
        used,
        [8, 26, 28, 14, 10, 14],
    )

    request_rows = [
        [
            request.id,
            request.customer.name if request.customer else "Unknown",
            request.service,
            request.request_date,
            request.status,
            request.technician or "",
            request.notes or "",
        ]
        for request in requests
    ]
    _table_sheet(
        workbook,
        "Service Requests",
        ["ID", "CUSTOMER", "SERVICE", "DATE", "STATUS", "TECHNICIAN", "NOTES"],
        request_rows,
        used,
        [8, 26, 28, 14, 14, 20, 42],
    )

    history_rows = [
        [
            record.id,
            record.customer.name if record.customer else "Unknown",
            record.service,
            record.service_date,
            record.technician or "",
            record.status,
            record.notes or "",
        ]
        for record in history
    ]
    _table_sheet(
        workbook,
        "Service History",
        ["ID", "CUSTOMER", "SERVICE", "DATE", "TECHNICIAN", "STATUS", "NOTES"],
        history_rows,
        used,
        [8, 26, 28, 14, 20, 14, 42],
    )

    concern_rows = [
        [
            concern.id,
            concern.customer.name if concern.customer else "Unknown",
            concern.concern,
            concern.status,
            concern.created_at,
        ]
        for concern in concerns
    ]
    _table_sheet(
        workbook,
        "Feedback",
        ["ID", "CUSTOMER", "FEEDBACK / CONCERN", "STATUS", "DATE"],
        concern_rows,
        used,
        [8, 26, 52, 14, 18],
    )

    notification_rows = [
        [
            notification.id,
            notification.customer.name if notification.customer else "Unknown",
            notification.message,
            notification.notification_date,
            notification.status,
        ]
        for notification in notifications
    ]
    _table_sheet(
        workbook,
        "Notifications",
        ["ID", "CUSTOMER", "MESSAGE", "DATE", "STATUS"],
        notification_rows,
        used,
        [8, 26, 52, 18, 14],
    )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
